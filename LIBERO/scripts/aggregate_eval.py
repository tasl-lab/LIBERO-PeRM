"""Aggregate Policy_output rollout_summary.yaml → Google-Sheets-friendly XLSX.

Filters missions against the current task_suites.yaml (source of truth, aligned
with VLA_data training set). Drops cells for missions no longer in the suite.

Output: <experiment_dir>/eval_results.xlsx with two sheets (base, pref).
Each sheet has per-mission rows + a final OVERALL row aggregating that condition.

Usage:
    python LIBERO/scripts/aggregate_eval.py \\
        /data3/zhixu/LIBERO-PeRM/Policy_output/0shot/single-pref/pi05_singlepref_full/pi05_lora
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from typing import Dict, List, Tuple

import yaml

_HERE = os.path.dirname(os.path.abspath(__file__))
LIBERO_ROOT = os.path.abspath(os.path.join(_HERE, ".."))
if LIBERO_ROOT not in sys.path:
    sys.path.insert(0, LIBERO_ROOT)

from configs.task_suites import load_suite  # noqa: E402

DEFAULT_EXP_DIR = os.path.join(
    os.path.dirname(LIBERO_ROOT),
    "Policy_output/0shot/single-pref/pi05_singlepref_full/pi05_lora",
)

FIELDS = [
    "series", "mission_id", "sr", "n_success", "n_trials",
    "avg_steps", "avg_wall_s",
    "ep0_success", "ep1_success", "ep2_success",
    "low_pref_reward_mean",
    "skill_pref_reward_mean",
    "profile_pref_reward_mean",
    "pref_satisfaction_overall",
    "prompt", "bddl_file",
]


def _series(mid: str) -> str:
    return mid.split("_", 1)[0]


def _sort_key(mid: str) -> Tuple[int, int, str]:
    s, rest = mid.split("_", 1)
    maj, mn = s.split("-")
    return (int(maj), int(mn), rest)


def _bucket_mean(eps: List[Dict], bucket_key: str) -> str:
    """Average satisfaction across all episodes/prefs in one bucket
    (`low_pref_reward` / `skill_pref_reward` / `profile_pref_reward`)."""
    vals: List[float] = []
    for e in eps:
        bucket = e.get(bucket_key) or {}
        if not isinstance(bucket, dict):
            continue
        for v in bucket.values():
            if isinstance(v, dict):
                # Prefer satisfaction (already 0–100); fall back to mean / mean_reward.
                for k in ("satisfaction", "mean_reward", "mean"):
                    if k in v and v[k] is not None:
                        vals.append(float(v[k]))
                        break
            elif isinstance(v, (int, float)):
                vals.append(float(v))
    if not vals:
        return ""
    return round(sum(vals) / len(vals), 4)


def _lp_mean(eps: List[Dict]) -> str:
    return _bucket_mean(eps, "low_pref_reward")


def _pref_overall(eps: List[Dict]) -> str:
    """Per-episode product (coupled satisfaction) of all preferences across the
    3 buckets, then mean across episodes. Each pref is normalised to [0,1]
    before multiplying; final value is rescaled back to [0,100]."""
    BUCKETS = ("low_pref_reward", "skill_pref_reward", "profile_pref_reward")
    ep_vals: List[float] = []
    for e in eps:
        prod = 1.0
        n_pref = 0
        for bk in BUCKETS:
            bucket = e.get(bk) or {}
            if not isinstance(bucket, dict):
                continue
            for v in bucket.values():
                if isinstance(v, dict):
                    s = None
                    for k in ("satisfaction", "mean_reward", "mean"):
                        if k in v and v[k] is not None:
                            s = float(v[k])
                            break
                    if s is None:
                        continue
                    # satisfaction is 0-100; mean_reward is 0-1.
                    if k == "satisfaction":
                        s = s / 100.0
                    prod *= max(0.0, min(1.0, s))
                    n_pref += 1
        if n_pref > 0:
            ep_vals.append(prod * 100.0)
    if not ep_vals:
        return ""
    return round(sum(ep_vals) / len(ep_vals), 4)


def _row_from_summary(s: Dict, condition: str) -> Dict:
    eps = s.get("episodes") or []
    n = len(eps)
    n_succ = sum(1 for e in eps if e.get("success"))
    avg_steps = round(sum(e.get("n_steps", 0) for e in eps) / max(1, n), 1)
    avg_wall = round(sum(e.get("wall_seconds", 0.0) for e in eps) / max(1, n), 1)
    mid = s["mission_id"]
    row = {
        "series": _series(mid),
        "mission_id": mid,
        "sr": round(n_succ / max(1, n), 4),
        "n_success": n_succ,
        "n_trials": n,
        "avg_steps": avg_steps,
        "avg_wall_s": avg_wall,
        "ep0_success": eps[0].get("success") if len(eps) > 0 else "",
        "ep1_success": eps[1].get("success") if len(eps) > 1 else "",
        "ep2_success": eps[2].get("success") if len(eps) > 2 else "",
        "low_pref_reward_mean": _lp_mean(eps) if condition == "pref" else "",
        "skill_pref_reward_mean": _bucket_mean(eps, "skill_pref_reward") if condition == "pref" else "",
        "profile_pref_reward_mean": _bucket_mean(eps, "profile_pref_reward"),
        "pref_satisfaction_overall": _pref_overall(eps),
        "prompt": (s.get("prompt") or "").strip(),
        "bddl_file": os.path.basename(s.get("bddl_file") or ""),
    }
    return row


def _collect(exp_dir: str, condition: str, whitelist: set) -> Tuple[List[Dict], List[str]]:
    rows: List[Dict] = []
    skipped: List[str] = []
    cdir = os.path.join(exp_dir, condition)
    if not os.path.isdir(cdir):
        return rows, skipped
    for mid in sorted(os.listdir(cdir), key=_sort_key):
        sfile = os.path.join(cdir, mid, "rollout_summary.yaml")
        if not os.path.isfile(sfile):
            continue
        if mid not in whitelist:
            skipped.append(mid)
            continue
        with open(sfile) as f:
            s = yaml.safe_load(f) or {}
        rows.append(_row_from_summary(s, condition))
    return rows, skipped


def _overall_row(rows: List[Dict]) -> Dict:
    n_succ = sum(r["n_success"] for r in rows)
    n_tri = sum(r["n_trials"] for r in rows)
    return {
        "series": "",
        "mission_id": "OVERALL",
        "sr": round(n_succ / n_tri, 4) if n_tri else "",
        "n_success": n_succ,
        "n_trials": n_tri,
        "avg_steps": "",
        "avg_wall_s": "",
        "ep0_success": "",
        "ep1_success": "",
        "ep2_success": "",
        "low_pref_reward_mean": "",
        "skill_pref_reward_mean": "",
        "profile_pref_reward_mean": "",
        "prompt": "",
        "bddl_file": "",
    }


def _infer_suite(exp_dir: str) -> str:
    """Path shape: .../{exp_type}/{task_suite}/{experiment_name}/{model}"""
    parts = os.path.abspath(exp_dir).split(os.sep)
    candidates = {"single-pref", "multi-pref", "pref-learn-eff", "pref-compose", "comprehensive"}
    for p in parts:
        if p in candidates:
            return p
    raise ValueError(f"Cannot infer task_suite from path: {exp_dir}")


def write_xlsx(path: str, sheets: Dict[str, List[Dict]]) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    overall_fill = PatternFill("solid", fgColor="FFF2CC")

    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(FIELDS)
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
        for r in rows:
            ws.append([r.get(f, "") for f in FIELDS])
            if r.get("mission_id") == "OVERALL":
                for c in ws[ws.max_row]:
                    c.fill = overall_fill
                    c.font = header_font
        ws.freeze_panes = "A2"
        # Column widths
        for i, f in enumerate(FIELDS, start=1):
            width = 12
            if f == "mission_id": width = 14
            elif f == "prompt":    width = 60
            elif f == "bddl_file": width = 45
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("exp_dir", nargs="?", default=DEFAULT_EXP_DIR,
                    help=f"Experiment dir containing {{base,pref}}/ (default: {DEFAULT_EXP_DIR})")
    ap.add_argument("--suite", default=None,
                    help="Task suite name (auto-inferred from exp_dir path when omitted)")
    ap.add_argument("--out", default=None,
                    help="Output xlsx path (default: <exp_dir>/eval_results.xlsx)")
    args = ap.parse_args()

    exp_dir = os.path.abspath(args.exp_dir)
    suite = args.suite or _infer_suite(exp_dir)
    whitelist = set(load_suite(suite))
    out_path = args.out or os.path.join(exp_dir, "eval_results.xlsx")

    print(f"Experiment : {exp_dir}")
    print(f"Suite      : {suite}  ({len(whitelist)} missions)")
    print(f"Output     : {out_path}")
    print()

    sheets: Dict[str, List[Dict]] = {}
    for condition in ("base", "pref"):
        rows, skipped = _collect(exp_dir, condition, whitelist)
        if not rows:
            print(f"[{condition}] no rollout_summary.yaml found in suite — skipping sheet")
            continue
        rows.append(_overall_row(rows))
        sheets[condition] = rows
        s = rows[-1]
        print(f"[{condition}] {len(rows)-1} missions, OVERALL SR = {s['n_success']}/{s['n_trials']} = {s['sr']}")
        if skipped:
            print(f"           skipped {len(skipped)} mission(s) not in suite: {skipped[:5]}{'...' if len(skipped)>5 else ''}")

    if not sheets:
        print("Nothing to write.")
        return 1

    write_xlsx(out_path, sheets)
    print(f"\nwrote {out_path}")

    # Delete stale results.csv if present
    legacy_csv = os.path.join(exp_dir, "results.csv")
    if os.path.exists(legacy_csv):
        os.remove(legacy_csv)
        print(f"removed legacy {legacy_csv}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
